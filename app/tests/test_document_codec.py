"""T045: the document codec channel (`cp-document`) and its worker engine.

The engine (`app.workers.document_service`) runs over the real frame codec and the
authenticated handshake of the fixed `cp-document` spec (the in-thread harness relaxes
only SO_PEERCRED, like the credential gateway's): PDF pages come back as bounded PNGs
checked by reopening, DOCX as bounded body/table text; every refusal is a closed code;
the control-side client verifies digests and the PNG size it can check without a decode.
The import-boundary tests pin that the control plane's artifact path never imports a PDF
or DOCX parser and never calls one while previewing through the worker.
"""

import ast
import io
import subprocess
import sys
import zipfile
from hashlib import sha256
from pathlib import Path

import pytest

from app.tests.support.document_worker_harness import in_thread_client
from app.tests.test_run_artifact_previews import docx_bytes, pdf_bytes
from app.workers import document_channel
from app.workers.document_channel import (
    MAX_EDGE_PX,
    MAX_INPUT_BYTES,
    MAX_TEXT_BYTES,
    DocumentCodecClient,
    DocumentCodecError,
    DocumentWorkerConfiguration,
)

REPOSITORY = Path(__file__).resolve().parents[2]
PARSERS = ("pypdfium2", "pypdfium2_raw", "pypdf", "PIL", "reportlab", "docx", "lxml")


@pytest.fixture
def harness():
    client, threads, outcomes = in_thread_client()
    yield client, outcomes
    for thread in threads:
        thread.join(5)
        assert not thread.is_alive()


def test_the_profile_is_the_compose_document_service():
    root, spec = document_channel.document_channel()
    assert str(root.pair_root) == "/run/deeptwin/ipc/cp-document"
    assert (spec.responder_uid, spec.responder_gid, spec.pair_gid) == (20106, 20106, 21105)
    assert (spec.requester_uid, spec.requester_gid) == (20102, 20102)
    assert spec.requester_message_types == ("document_input", "document_request")
    assert spec.responder_message_types == ("document_output", "document_result")
    compose = (REPOSITORY / "deploy" / "compose.yaml").read_text()
    assert '"user": "20106:20106"' in compose and '"network_mode": "none"' in compose
    assert "--channel=cp-document" in compose and "--endpoint-owner=20106:21105" in compose


def test_a_pdf_page_is_rendered_bounded_and_reported_by_the_real_page_count(harness):
    client, outcomes = harness
    pdf = pdf_bytes(3)
    page = client.render_page(pdf, page=2, max_edge_px=640)
    assert (page.page, page.page_count) == (2, 3)
    assert max(page.width, page.height) <= 640 and page.height == 640  # portrait letter
    assert (page.page_width_pt, page.page_height_pt) == (612, 792)
    assert page.png.startswith(b"\x89PNG") and sha256(page.png).hexdigest() == page.sha256
    small = client.render_page(pdf, page=1, max_edge_px=64)
    assert max(small.width, small.height) == 64
    with pytest.raises(DocumentCodecError) as missing:
        client.render_page(pdf, page=4)
    assert missing.value.code == "page_out_of_range" and missing.value.page_count == 3
    assert outcomes == ["ok", "ok", "page_out_of_range"]


@pytest.mark.parametrize("data, code", [
    (b"%PDF-1.7\n1 0 obj << >> endobj trailer <<>>", "content_rejected"),
    (b"not a pdf at all", "content_rejected"),
])
def test_damaged_pdfs_are_refused_not_rendered(harness, data, code):
    client, _outcomes = harness
    with pytest.raises(DocumentCodecError) as refused:
        client.render_page(data, page=1)
    assert refused.value.code == code


def test_an_encrypted_pdf_is_refused(harness):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(pdf_bytes(1))).pages:
        writer.add_page(page)
    writer.encrypt("owner-secret-only")
    buffer = io.BytesIO()
    writer.write(buffer)
    client, _outcomes = harness
    with pytest.raises(DocumentCodecError) as refused:
        client.render_page(buffer.getvalue(), page=1)
    assert refused.value.code == "content_rejected"


def test_docx_text_is_extracted_in_document_order_and_bounded(harness):
    client, _outcomes = harness
    text = client.extract_text(docx_bytes())
    assert text.text == "보고서 첫 문단\n항목\t값\n가\t1\n<script>alert(1)</script>"
    assert (text.paragraphs, text.tables, text.table_cells, text.truncated) == (2, 1, 4, False)
    assert text.sha256 == sha256(text.text.encode()).hexdigest()
    from docx import Document

    long = Document()
    for _ in range(3000):
        long.add_paragraph("가나다라마바사아자차카타파하" * 2)
    buffer = io.BytesIO()
    long.save(buffer)
    cut = client.extract_text(buffer.getvalue())
    assert cut.truncated is True and len(cut.text.encode()) <= MAX_TEXT_BYTES
    cut.text.encode("utf-8")  # cut on a character boundary


def _zip(entries):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return buffer.getvalue()


def test_containers_that_are_not_docx_or_are_bombs_are_refused(harness):
    client, _outcomes = harness
    with pytest.raises(DocumentCodecError) as other:
        client.extract_text(_zip([("xl/workbook.xml", "<x/>")]))
    assert other.value.code == "media_unsupported"
    with pytest.raises(DocumentCodecError) as bomb:
        client.extract_text(_zip([("word/document.xml", b"\x00" * (70 * 1024 * 1024))]))
    assert bomb.value.code == "content_rejected"
    with pytest.raises(DocumentCodecError) as garbage:
        client.extract_text(b"PK\x03\x04garbage")
    assert garbage.value.code == "content_rejected"


def test_the_client_bounds_requests_before_anything_is_sent():
    sent = []
    client = DocumentCodecClient(lambda: sent.append(1))
    for call in (lambda: client.render_page(b"x", page=0), lambda: client.render_page(b"x", page=1, max_edge_px=MAX_EDGE_PX + 1),
                 lambda: client.render_page(b"", page=1), lambda: client.render_page(b"x" * (MAX_INPUT_BYTES + 1), page=1)):
        with pytest.raises(DocumentCodecError) as refused:
            call()
        assert refused.value.sent is False
    assert sent == []


class _LyingService:
    """A worker that answers with a result whose declared size is not its PNG's."""

    def __init__(self, mode):
        self.mode = mode

    def serve_connection(self, connection, *, deadline):
        from uuid import uuid4

        from app.workers.document_channel import (
            RESULT_SCHEMA, RESULT_TYPE, canonical, chunk_count, read_chunks, strict_object, write_chunks)

        first = connection.read(deadline=deadline)
        request = strict_object(first.payload)
        read_chunks(connection, size=request["input_bytes"], count=request["chunk_count"],
                    digest=request["input_sha256"], message_type="document_input",
                    correlation_id=first.envelope.message_id, deadline=deadline)
        png = pdf_page_png()
        digest = sha256(png).hexdigest() if self.mode != "digest" else "0" * 64
        result = {"page": 1, "page_count": 1, "width": 10 if self.mode == "size" else 1,
                  "height": 1, "page_width_pt": 1, "page_height_pt": 1}
        if self.mode == "code":
            header = {"schema": RESULT_SCHEMA, "ok": False, "code": "rm -rf"}
        else:
            header = {"schema": RESULT_SCHEMA, "ok": True, "result": result, "output_bytes": len(png),
                      "output_sha256": digest, "chunk_count": chunk_count(len(png))}
        connection.write(message_id=str(uuid4()), correlation_id=first.envelope.message_id,
                         message_type=RESULT_TYPE, payload=canonical(header), deadline=deadline)
        if self.mode != "code":
            write_chunks(connection, png, message_type="document_output",
                         correlation_id=first.envelope.message_id, deadline=deadline)
        return "lied"


def pdf_page_png():
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (1, 1)).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.mark.parametrize("mode", ["size", "digest", "code"])
def test_the_client_refuses_a_worker_result_it_can_check(mode):
    client, threads, _outcomes = in_thread_client(service=_LyingService(mode), deadline_ms=5_000)
    with pytest.raises(DocumentCodecError) as refused:
        client.render_page(b"%PDF-1.7", page=1)
    assert refused.value.code == "malformed_result"
    for thread in threads:
        thread.join(5)


def test_a_failed_connection_is_unavailable_and_unsent():
    def refuse():
        raise OSError("no worker")

    with pytest.raises(DocumentCodecError) as refused:
        DocumentCodecClient(refuse).render_page(b"%PDF-1.7", page=1)
    assert refused.value.code == "unavailable" and refused.value.sent is False


def test_the_attachment_names_only_the_fixed_pair_root():
    value = {"schema": document_channel.ATTACHMENT_SCHEMA, "pair_root": "/run/deeptwin/ipc/cp-document",
             "requester_boot_id": "control-boot"}
    configuration = DocumentWorkerConfiguration.from_mapping(value)
    assert type(DocumentCodecClient.for_worker(configuration)) is DocumentCodecClient
    for bad in ({**value, "schema": "x"}, {**value, "extra": 1}, {**value, "pair_root": "relative"},
                {**value, "requester_boot_id": "-bad boot"}):
        with pytest.raises(ValueError):
            DocumentWorkerConfiguration.from_mapping(bad)
    with pytest.raises(ValueError):
        DocumentCodecClient.for_worker(DocumentWorkerConfiguration(pair_root="/tmp/cp-document",
                                                                   requester_boot_id="control-boot"))


# --- the import boundary ---------------------------------------------------------------

CONTROL_PATH = ("app/services/run_artifacts.py", "app/workers/document_channel.py", "app/api/artifact_index.py",
                "app/api/runs.py", "app/api/first_party_catalog.py", "app/server.py")


def test_the_control_plane_artifact_path_imports_no_document_parser():
    for relative in CONTROL_PATH:
        tree = ast.parse((REPOSITORY / relative).read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            names = ([alias.name for alias in node.names] if isinstance(node, ast.Import)
                     else [node.module or ""] if isinstance(node, ast.ImportFrom) else [])
            for name in names:
                assert name.split(".")[0] not in PARSERS, (relative, name)
                assert not name.endswith("document_service"), (relative, name)


def test_the_control_plane_process_never_loads_a_pdf_parser():
    code = ("import sys, app.server, app.api.first_party_catalog, app.services.run_artifacts; "
            "print(sorted({n.split('.')[0] for n in sys.modules if n.split('.')[0] in "
            "('pypdfium2', 'pypdfium2_raw', 'pypdf', 'PIL', 'reportlab')}), "
            "'app.workers.document_service' in sys.modules)")
    result = subprocess.run([sys.executable, "-B", "-c", code], cwd=REPOSITORY, capture_output=True,
                            text=True, timeout=120, check=True)
    assert result.stdout.strip() == "[] False"


def test_the_worker_entrypoint_imports_no_web_or_control_plane_code():
    code = ("import sys, app.workers.document_worker_main as m; "
            "print(sorted(n for n in sys.modules if n.split('.')[0] in ('fastapi', 'starlette', 'uvicorn')"
            " or n.startswith(('app.api', 'app.server', 'app.services'))))")
    result = subprocess.run([sys.executable, "-B", "-c", code], cwd=REPOSITORY, capture_output=True,
                            text=True, timeout=60, check=True)
    assert result.stdout.strip() == "[]"


def test_previews_through_the_worker_never_call_a_parser_in_this_process(tmp_path, monkeypatch):
    """The legacy upload extractor (`app.ingestion`, loaded by the domain store) imports
    python-docx into every control-plane process; the artifact path must still never call
    it. With every parser entry point in this process poisoned, the worker process still
    answers both previews."""

    import secrets
    import shutil
    import tempfile

    import docx

    from app.tests.support.document_worker_harness import process_client
    from app.tests.test_run_artifact_previews import DOCX_MEDIA, get, owner_app, started

    pdf, document = pdf_bytes(2), docx_bytes()

    def poisoned(*_args, **_kwargs):
        raise AssertionError("a parser was called in the control plane")

    directory = Path(tempfile.mkdtemp(prefix="dt-t045-boundary-"))
    path = directory / "document.sock"
    secret = secrets.token_bytes(32)
    process = subprocess.Popen([sys.executable, "-B", "-m", "app.tests.support.document_worker_harness", str(path)],
                               cwd=REPOSITORY, stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    process.stdin.write(secret.hex().encode() + b"\n")
    process.stdin.close()
    try:
        assert process.stdout.readline().strip() == b"ready"
        for name in ("pypdfium2", "pypdfium2_raw", "pypdf", "PIL", "PIL.Image"):
            monkeypatch.setitem(sys.modules, name, None)
        monkeypatch.setattr(docx, "Document", poisoned)
        monkeypatch.setattr(docx.api, "Document", poisoned)
        import app.ingestion

        monkeypatch.setattr(app.ingestion, "Document", poisoned)
        with owner_app(tmp_path, document_worker=process_client(path, secret)) as subject:
            from app.services.run_artifacts import artifact_identity

            run_id, ref = started(subject, [("report", "application/pdf", pdf), ("draft", DOCX_MEDIA, document)])
            page = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 0)}/pages/2").json()["preview"]
            assert page["kind"] == "page_image" and page["body"]["page_count"] == 2
            text = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 1)}/preview").json()["preview"]
            assert text["kind"] == "document_text"
    finally:
        process.kill()
        process.wait(5)
        shutil.rmtree(directory, ignore_errors=True)
