from io import BytesIO
import shutil
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pytest

from app.ingestion import extract


@pytest.mark.parametrize('name', ['a.txt', 'a.md', 'a.csv', 'a.json'])
def test_plain_files_preserve_literal_utf8(name):
    body = '\n한국어 <script>do not execute</script>\n=FORMULA()'
    result = extract(name, body.encode())
    assert result['read_status'] == 'read' and result['text'] == body
    assert set(result) == {'read_status', 'text', 'read_note', 'media_type'}


def test_bad_encoding_and_length_are_honest():
    assert extract('a.txt', b'\xff')['read_status'] == 'unreadable'
    result = extract('a.txt', b'a' * 20_001)
    assert result['read_status'] == 'partial' and len(result['text']) == 20_000
    assert result['read_note']


def test_docx_paragraphs_and_tables_are_read_without_execution():
    doc = Document()
    doc.add_paragraph('문단 원문')
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = '표 내용'
    stream = BytesIO()
    doc.save(stream)
    result = extract('source.docx', stream.getvalue())
    assert result['read_status'] == 'partial'
    assert '문단 원문' in result['text'] and '표 내용' in result['text']
    assert result['read_note']


def test_corrupt_and_oversized_docx_are_not_parsed():
    assert extract('bad.docx', b'not zip')['read_status'] == 'unreadable'
    stream = BytesIO()
    with ZipFile(stream, 'w') as archive:
        archive.writestr('../escape.xml', 'outside')
    assert extract('bad.docx', stream.getvalue())['read_status'] == 'unreadable'


@pytest.mark.parametrize('name', ['a.png', 'a.jpg', 'a.jpeg', 'a.svg', 'a.html'])
def test_unprocessed_images_and_active_documents_are_not_claimed_read(name):
    result = extract(name, b'<script>literal untrusted file</script>')
    assert result['read_status'] == 'unsupported' and result['text'] == ''
    assert result['read_note']


def test_unknown_format_is_rejected():
    with pytest.raises(ValueError):
        extract('command.sh', b'echo bad')


def test_existing_synthetic_pdf_has_honest_text_only_status():
    asset = next((Path(__file__).parents[2] / 'control-prototype/assets').glob('*.pdf'))
    result = extract(asset.name, asset.read_bytes())
    if shutil.which('pdftotext'):
        assert result['read_status'] == 'partial' and result['text'].strip()
    else:
        assert result['read_status'] == 'unsupported' and not result['text']
    assert result['read_note']


def test_pdf_process_output_is_bounded_and_marked_partial(tmp_path, monkeypatch):
    # Exercise a real child process without depending on installed Poppler.
    executable = tmp_path / 'pdftotext'
    executable.write_text(f'#!{sys.executable}\nimport sys\nsys.stdout.write("x" * 1_000_000)\n')
    executable.chmod(0o700)
    monkeypatch.setattr(shutil, 'which', lambda name: str(executable))
    asset = next((Path(__file__).parents[2] / 'control-prototype/assets').glob('*.pdf'))
    result = extract(asset.name, asset.read_bytes())
    assert result['read_status'] == 'partial'
    assert len(result['text']) == 20_000


def test_docx_decompression_limit_is_checked_before_parsing():
    stream = BytesIO()
    with ZipFile(stream, 'w') as archive:
        archive.writestr('word/document.xml', b'x' * (20 * 1024 * 1024 + 1))
    with pytest.raises(ValueError):
        extract('large.docx', stream.getvalue())


def test_pdf_without_extractable_text_is_not_marked_partly_read(tmp_path, monkeypatch):
    executable = tmp_path / 'pdftotext'
    executable.write_text(f'#!{sys.executable}\nprint("  ")\n')
    executable.chmod(0o700)
    monkeypatch.setattr(shutil, 'which', lambda name: str(executable))
    asset = next((Path(__file__).parents[2] / 'control-prototype/assets').glob('*.pdf'))
    result = extract(asset.name, asset.read_bytes())
    assert result['read_status'] == 'unreadable' and result['text'] == ''
    assert 'OCR' in result['read_note']
