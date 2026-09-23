"""US3 bounded declarative document creation with real render inspection.

DOCX, CSV and JSON creation is declarative and bounded; every render is
verified by REOPENING the produced file and inspecting its actual content —
a file's existence is never the evidence (SC-004). Formula-looking CSV
values remain data: the default writer preserves them verbatim, and the
optional safe-spreadsheet export records every escaping transformation it
applied. Format validation sniffs real bytes, not extensions: JSON bytes
declared as DOCX refuse, NUL-poisoned text refuses, and a plain ZIP is not
a DOCX (runtime.md §6 artifact tools, T044). PDFs are read back through
their text layer and rasterized; PNGs are reopened and sampled.
"""

import json

import pytest

from app.adapters.documents import (
    DocumentToolError,
    render_csv,
    render_docx,
    render_json,
    render_pdf,
    render_png,
    validate_format,
)


def docx_spec(**overrides):
    value = {
        "title": "주간 대본 보고",
        "paragraphs": [
            "research 노드가 수집한 출처 요약이다.",
            "writer 노드의 최종 대본 개요이다.",
        ],
        "tables": [
            {
                "rows": [
                    ["항목", "값"],
                    ["출처 수", "12"],
                ],
            },
        ],
    }
    value.update(overrides)
    return value


def test_docx_renders_and_is_inspected_by_reopening(tmp_path):
    path = tmp_path / "report.docx"
    report = render_docx(docx_spec(), path)
    # the report's facts come from REOPENING the produced file
    assert report.paragraph_count >= 3  # title + two paragraphs
    assert report.table_count == 1
    assert report.inspected_cells[("t0", 1, 1)] == "12"
    assert report.byte_length == path.stat().st_size
    assert len(report.sha256) == 64
    from docx import Document

    reopened = Document(str(path))
    texts = [p.text for p in reopened.paragraphs]
    assert "research 노드가 수집한 출처 요약이다." in texts


def test_docx_bounds_are_enforced(tmp_path):
    with pytest.raises(DocumentToolError):
        render_docx(docx_spec(paragraphs=["x"] * 501), tmp_path / "a.docx")
    with pytest.raises(DocumentToolError):
        render_docx(docx_spec(
            tables=[{"rows": [["c"] * 51]}],
        ), tmp_path / "b.docx")
    with pytest.raises(DocumentToolError):
        render_docx(docx_spec(paragraphs=["y" * 10_000]), tmp_path / "c.docx")
    with pytest.raises(DocumentToolError):
        render_docx({"unexpected": True}, tmp_path / "d.docx")


def test_csv_keeps_formula_looking_values_as_data(tmp_path):
    path = tmp_path / "table.csv"
    rows = [["name", "total"], ["=SUM(A1:A9)", "+82"], ["@cmd", "-3"]]
    report = render_csv(rows, path)
    assert report.transformations == ()  # data stays data by default
    import csv as _csv

    with open(path, newline="", encoding="utf-8") as handle:
        read_back = list(_csv.reader(handle))
    assert read_back == rows  # verified by reopening, byte-faithful


def test_safe_spreadsheet_export_records_every_transformation(tmp_path):
    path = tmp_path / "safe.csv"
    rows = [["=SUM(A1)", "plain"], ["ok", "@risky"]]
    report = render_csv(rows, path, safe_spreadsheet=True)
    assert report.transformations == (
        (0, 0, "prefixed_apostrophe"),
        (1, 1, "prefixed_apostrophe"),
    )
    import csv as _csv

    with open(path, newline="", encoding="utf-8") as handle:
        read_back = list(_csv.reader(handle))
    assert read_back[0][0] == "'=SUM(A1)"
    assert read_back[0][1] == "plain"


def test_json_round_trips_and_bounds(tmp_path):
    path = tmp_path / "value.json"
    value = {"이름": "DeepTwin", "숫자": [1, 2, 3]}
    report = render_json(value, path)
    assert report.byte_length == path.stat().st_size
    assert json.loads(path.read_text(encoding="utf-8")) == value
    with pytest.raises(DocumentToolError):
        render_json({"nan": float("nan")}, tmp_path / "bad.json")
    deep = value
    for _ in range(100):
        deep = [deep]
    with pytest.raises(DocumentToolError):
        render_json(deep, tmp_path / "deep.json")


def test_format_validation_sniffs_real_bytes(tmp_path):
    docx_path = tmp_path / "real.docx"
    render_docx(docx_spec(), docx_path)
    assert validate_format(docx_path, "docx") is True
    json_path = tmp_path / "fake.docx"
    json_path.write_text('{"not": "a docx"}', encoding="utf-8")
    with pytest.raises(DocumentToolError):
        validate_format(json_path, "docx")
    import zipfile

    plain_zip = tmp_path / "plain.docx"
    with zipfile.ZipFile(plain_zip, "w") as archive:
        archive.writestr("hello.txt", "not a document")
    with pytest.raises(DocumentToolError):
        validate_format(plain_zip, "docx")  # a plain zip is not a DOCX
    nul_csv = tmp_path / "evil.csv"
    nul_csv.write_bytes(b"a,b\x00c\n")
    with pytest.raises(DocumentToolError):
        validate_format(nul_csv, "csv")
    good_csv = tmp_path / "good.csv"
    render_csv([["a", "b"]], good_csv)
    assert validate_format(good_csv, "csv") is True
    with pytest.raises(DocumentToolError):
        validate_format(docx_path, "json")
    with pytest.raises(DocumentToolError):
        validate_format(docx_path, "pdf")  # a DOCX is not a PDF
    with pytest.raises(DocumentToolError):
        validate_format(docx_path, "xlsx")  # unsupported here stays explicit


# --- PDF and PNG (T044) -------------------------------------------------------


def text_on_page(path):
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        return "".join(document[i].get_textpage().get_text_range() for i in range(len(document)))
    finally:
        document.close()


def test_pdf_renders_hangul_and_tables_and_is_inspected_by_rasterizing(tmp_path):
    path = tmp_path / "report.pdf"
    report = render_pdf(docx_spec(), path)
    assert report.media == "pdf" and report.title_verified
    assert report.table_count == 1 and report.inspected_cells[("t0", 1, 1)] == "12"
    assert report.byte_length == path.stat().st_size
    text = "".join(text_on_page(path).split())
    for item in ("주간대본보고", "research노드가수집한출처요약이다.", "출처수", "12"):
        assert item in text
    assert validate_format(path, "pdf") is True
    # invariant output: the same spec renders the same bytes
    again = render_pdf(docx_spec(), tmp_path / "again.pdf")
    assert again.sha256 == report.sha256


def test_declared_pdf_text_is_data_never_markup(tmp_path):
    path = tmp_path / "markup.pdf"
    render_pdf(docx_spec(paragraphs=["<b>굵게</b> & <font color=red>x</font> &amp;"]), path)
    text = "".join(text_on_page(path).split())
    assert "<b>굵게</b>&<fontcolor=red>x</font>&amp;" in text


def test_a_long_document_wraps_across_pages_and_is_still_found_in_order(tmp_path):
    spec = docx_spec(paragraphs=[f"{index}번째 문단: " + "가나다라마바사 " * 30
                                 for index in range(60)])
    report = render_pdf(spec, tmp_path / "long.pdf")
    from pypdf import PdfReader

    assert len(PdfReader(str(tmp_path / "long.pdf")).pages) > 1
    assert report.paragraph_count == 60


def test_a_glyph_the_font_cannot_draw_fails_the_render(tmp_path, monkeypatch):
    import pypdfium2 as pdfium

    blank = pdfium.PdfPage.render

    def white(self, **kwargs):
        rendered = blank(self, **kwargs)
        image = rendered.to_pil()
        image.paste((255, 255, 255), (0, 0, image.width, image.height))

        class Bitmap:
            def to_pil(self):
                return image

        return Bitmap()

    monkeypatch.setattr(pdfium.PdfPage, "render", white)
    with pytest.raises(DocumentToolError, match="no ink"):
        render_pdf(docx_spec(), tmp_path / "blank.pdf")


def test_pdf_specs_are_bounded(tmp_path):
    for spec in (
        docx_spec(title=""),
        docx_spec(paragraphs=["a\x00b"]),
        docx_spec(tables=[{"rows": [["a"], ["b", "c"]]}]),
        {"title": "x"},
    ):
        with pytest.raises(DocumentToolError):
            render_pdf(spec, tmp_path / "bad.pdf")


@pytest.mark.parametrize("payload", [
    b"/OpenAction << /S /JavaScript /JS (app.alert(1)) >>",
    b"/AA << /O << /S /Launch /F (calc.exe) >> >>",
    b"/Names << /EmbeddedFiles 3 0 R >>",
    b"/URI (https://example.invalid/)",
])
def test_active_pdf_content_is_refused(tmp_path, payload):
    source = tmp_path / "clean.pdf"
    render_pdf(docx_spec(), source)
    data = source.read_bytes()
    hostile = tmp_path / "hostile.pdf"
    hostile.write_bytes(data.replace(b"/Type /Catalog", b"/Type /Catalog " + payload, 1))
    with pytest.raises(DocumentToolError, match="active content"):
        validate_format(hostile, "pdf")


def test_a_link_hidden_in_a_compressed_object_stream_is_refused(tmp_path):
    from pypdf import PdfWriter
    from pypdf.annotations import Link

    source = tmp_path / "clean.pdf"
    render_pdf(docx_spec(), source)
    writer = PdfWriter(clone_from=str(source))
    writer.add_annotation(0, Link(rect=(10, 10, 50, 50), url="https://example.invalid/"))
    writer.compress_identical_objects()
    hidden = tmp_path / "hidden.pdf"
    with open(hidden, "wb") as handle:
        writer.write(handle)
    with pytest.raises(DocumentToolError, match="active content"):
        validate_format(hidden, "pdf")
    assert validate_format(source, "pdf") is True


def test_corrupt_and_mislabelled_pdfs_are_refused(tmp_path):
    source = tmp_path / "clean.pdf"
    render_pdf(docx_spec(), source)
    truncated = tmp_path / "truncated.pdf"
    truncated.write_bytes(source.read_bytes()[:400])
    with pytest.raises(DocumentToolError):
        validate_format(truncated, "pdf")
    fake = tmp_path / "fake.pdf"
    fake.write_bytes(b'{"a": 1}')
    with pytest.raises(DocumentToolError):
        validate_format(fake, "pdf")


def image_spec(**overrides):
    value = {
        "width": 64, "height": 40, "background": [255, 255, 255],
        "shapes": [
            {"x": 0, "y": 0, "width": 32, "height": 40, "fill": [200, 30, 30]},
            {"x": 16, "y": 10, "width": 32, "height": 20, "fill": [30, 30, 200]},
        ],
    }
    value.update(overrides)
    return value


def test_png_renders_and_is_inspected_by_sampling_pixels(tmp_path):
    path = tmp_path / "chart.png"
    report = render_png(image_spec(), path)
    assert report.media == "png"
    # the overlapping second shape is what the first one's centre shows
    assert report.inspected_cells[("s0", 16, 20)] == (30, 30, 200)
    assert report.inspected_cells[("s1", 32, 20)] == (30, 30, 200)
    from PIL import Image

    with Image.open(path) as image:
        assert image.size == (64, 40) and image.getpixel((63, 0))[:3] == (255, 255, 255)
    assert validate_format(path, "png") is True


def test_png_specs_are_bounded(tmp_path):
    for spec in (
        image_spec(width=0),
        image_spec(height=5000),
        image_spec(background=[256, 0, 0]),
        image_spec(shapes=[{"x": 60, "y": 0, "width": 10, "height": 1, "fill": [0, 0, 0]}]),
        image_spec(shapes=[{"x": 0, "y": 0, "width": 1, "height": 1, "fill": [0, 0]}]),
        image_spec(shapes=[{"x": True, "y": 0, "width": 1, "height": 1, "fill": [0, 0, 0]}]),
    ):
        with pytest.raises(DocumentToolError):
            render_png(spec, tmp_path / "bad.png")


def test_corrupt_oversized_and_mislabelled_pngs_are_refused(tmp_path):
    good = tmp_path / "good.png"
    render_png(image_spec(), good)
    truncated = tmp_path / "truncated.png"
    truncated.write_bytes(good.read_bytes()[:60])
    with pytest.raises(DocumentToolError):
        validate_format(truncated, "png")
    from PIL import Image

    huge = tmp_path / "huge.png"
    Image.new("1", (5000, 10)).save(huge)
    with pytest.raises(DocumentToolError):
        validate_format(huge, "png")
    with pytest.raises(DocumentToolError):
        validate_format(good, "pdf")
