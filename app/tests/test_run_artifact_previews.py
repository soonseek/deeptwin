"""T045: PDF/DOCX previews through the isolated document worker, and the vault-wide
artifact index, read by the owner through the supported server.

The document service runs here as a SEPARATE PROCESS (`app.tests.support.
document_worker_harness`) serving the unmodified `DocumentCodecService` over the real
frame codec and authenticated handshake of the `cp-document` spec; the supported factory
is given the frame-only client for it (`create_app(..., document_worker=client)`). The
control plane never parses the PDF or the DOCX: the page arrives as the worker's PNG with
its digest, the text as the worker's extraction; a worker that is not named, cannot be
reached, refuses the file or drops the dialogue is disclosed, never replaced by a
stand-in. The root-only production entrypoint over the real pair root with SO_PEERCRED is
qualified separately (`test_document_worker_main.py`).
"""

import io
import secrets
import shutil
import subprocess
import sys
import tempfile
from contextlib import contextmanager
from hashlib import sha256
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from app.runtime.budgets import BudgetPolicy
from app.server import create_app
from app.services import run_artifacts
from app.services.run_artifacts import artifact_identity
from app.tests.support.document_worker_harness import in_thread_client, process_client
from app.tests.test_graph_execution import linear_graph
from app.tests.test_run_artifacts_api import PNG, sealed_result
from app.tests.test_runs_api import Executor, command, graph_record, post
from app.tests.test_server_api_v1 import immutable
from app.tests.test_web_owner_integration import bootstrap_client, configured, headers
from app.workers.document_channel import DocumentCodecClient

REPOSITORY = Path(__file__).resolve().parents[2]
DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def pdf_bytes(pages=3):
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    drawing = canvas.Canvas(buffer, pagesize=(612, 792))
    for number in range(pages):
        drawing.setFont("Helvetica", 48)
        drawing.drawString(100, 600, f"Page {number + 1}")
        drawing.rect(80, 560 - number * 40, 300, 20, fill=1)
        drawing.showPage()
    drawing.save()
    return buffer.getvalue()


def docx_bytes():
    from docx import Document

    document = Document()
    document.add_paragraph("보고서 첫 문단")
    table = document.add_table(rows=2, cols=2)
    for row, values in enumerate((("항목", "값"), ("가", "1"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    document.add_paragraph("<script>alert(1)</script>")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def worker_process():
    directory = Path(tempfile.mkdtemp(prefix="dt-t045-doc-"))
    path = directory / "document.sock"
    secret = secrets.token_bytes(32)
    process = subprocess.Popen([sys.executable, "-B", "-m", "app.tests.support.document_worker_harness", str(path)],
                               cwd=REPOSITORY, stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                               stderr=subprocess.PIPE)
    process.stdin.write(secret.hex().encode() + b"\n")
    process.stdin.close()
    assert process.stdout.readline().strip() == b"ready"
    try:
        yield SimpleNamespace(path=path, secret=secret, process=process)
    finally:
        process.kill()
        process.wait(5)
        shutil.rmtree(directory, ignore_errors=True)


@contextmanager
def owner_app(tmp_path, *, document_worker=None):
    tmp_path.mkdir(parents=True, exist_ok=True)
    executor = Executor()
    profile, capability, arguments = configured(tmp_path)
    app = create_app(tmp_path / "data", **arguments, run_executor=executor, document_worker=document_worker)
    with TestClient(app, base_url=profile.http_origin) as client:
        csrf = bootstrap_client(client, profile, capability)
        domain = app.state.domain_store
        roots = domain.roots()
        policy = BudgetPolicy.create(
            profile="execution", provider_mode="subscription", max_model_calls=3,
            max_tool_calls=5, max_node_visits=7, max_loop_rounds=2, max_output_bytes=1_000,
            max_concurrency=2, max_wall_seconds=60, max_candidates=1)
        refs = SimpleNamespace(
            work=immutable(domain, roots, "work_revision").ref,
            environment=immutable(domain, roots, "environment").ref,
            budget=immutable(domain, roots, "budget_policy", content=policy.domain_content()).ref)
        yield SimpleNamespace(app=app, client=client, profile=profile, csrf=csrf, refs=refs,
                              domain=domain, path=profile.base_path + "api/v1/runs", executor=executor)


def started(subject, outputs):
    subject.executor.result = sealed_result(subject, outputs)
    created = post(subject, command(subject, graph_record(subject, linear_graph())))
    assert created.status_code == 201, created.text
    return created.json()["run_id"], subject.executor.result


def get(subject, path, **extra):
    return subject.client.get(subject.path + path, headers={**headers(subject.profile), **extra})


def index(subject, query=""):
    return subject.client.get(subject.profile.base_path + "api/v1/artifacts" + query,
                              headers=headers(subject.profile))


def test_a_pdf_page_and_a_docx_text_come_from_the_separate_worker_process(tmp_path, worker_process):
    client = process_client(worker_process.path, worker_process.secret)
    pdf, docx = pdf_bytes(3), docx_bytes()
    with owner_app(tmp_path, document_worker=client) as subject:
        run_id, ref = started(subject, [("report", "application/pdf", pdf), ("draft", DOCX_MEDIA, docx)])
        report = artifact_identity(ref, 0)
        first = get(subject, f"/{run_id}/artifacts/{report}/preview")
        assert first.status_code == 200, first.text
        preview = first.json()["preview"]
        assert preview["kind"] == "page_image" and preview["derived"] is True
        assert preview["body"]["page"] == 1 and preview["body"]["page_count"] == 3
        assert max(preview["body"]["width"], preview["body"]["height"]) <= 1200
        assert preview["coverage"] == {"original_bytes": len(pdf), "unit": "page", "page_count": 3,
                                       "covered": [[1, 1]], "omissions": [[2, 3]],
                                       "not_rendered": ["annotations_interactive", "links", "text_layer",
                                                        "forms", "attachments"]}
        image = get(subject, f"/{run_id}/artifacts/{report}/pages/1/image")
        assert image.status_code == 200 and image.headers["content-type"] == "image/png"
        assert image.content.startswith(b"\x89PNG\r\n\x1a\n")
        assert sha256(image.content).hexdigest() == preview["sha256"] == image.headers["x-deeptwin-derived-sha256"]
        assert "default-src 'none'" in image.headers["content-security-policy"]
        assert image.headers["x-content-type-options"] == "nosniff"
        second = get(subject, f"/{run_id}/artifacts/{report}/pages/2")
        assert second.status_code == 200, second.text
        page = second.json()["preview"]
        assert page["body"]["page"] == 2 and page["coverage"]["omissions"] == [[1, 1], [3, 3]]
        assert page["sha256"] != preview["sha256"]  # a different page, a different derived image
        # the original is untouched and still served as-is
        original = get(subject, f"/{run_id}/artifacts/{report}/content")
        assert original.content == pdf and original.headers["content-type"] == "application/octet-stream"
        text = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 1)}/preview").json()["preview"]
        assert text["kind"] == "document_text" and text["derived"] is True
        assert text["body"]["text"] == "보고서 첫 문단\n항목\t값\n가\t1\n<script>alert(1)</script>"
        assert text["sha256"] == sha256(text["body"]["text"].encode()).hexdigest()
        assert (text["body"]["paragraphs"], text["body"]["tables"], text["body"]["table_cells"]) == (2, 1, 4)
        assert "headers_footers" in text["coverage"]["omissions"] and text["coverage"]["unit"] == "text_layer"


def test_page_bounds_media_and_size_are_refused_honestly(tmp_path, worker_process, monkeypatch):
    client = process_client(worker_process.path, worker_process.secret)
    pdf = pdf_bytes(2)
    with owner_app(tmp_path, document_worker=client) as subject:
        run_id, ref = started(subject, [("report", "application/pdf", pdf), ("note", "text/plain", b"hello"),
                                        ("broken", "application/pdf", b"%PDF-1.7\nnot really a pdf")])
        report, note, broken = (artifact_identity(ref, index) for index in range(3))
        out_of_range = get(subject, f"/{run_id}/artifacts/{report}/pages/9")
        assert out_of_range.status_code == 416 and out_of_range.json()["code"] == "range_not_satisfiable"
        assert get(subject, f"/{run_id}/artifacts/{report}/pages/9/image").status_code == 416
        for bad in ("0", "01", "x", "-1", "123456"):
            assert get(subject, f"/{run_id}/artifacts/{report}/pages/{bad}").status_code == 400
        assert get(subject, f"/{run_id}/artifacts/{report}/pages/1/thumb").status_code == 404
        not_a_pdf = get(subject, f"/{run_id}/artifacts/{note}/pages/1")
        assert not_a_pdf.status_code == 415 and not_a_pdf.json()["code"] == "unsupported_media"
        # the worker refuses a damaged file: disclosed as a worker failure, never a page
        failed = get(subject, f"/{run_id}/artifacts/{broken}/pages/1")
        assert failed.status_code == 502 and failed.json()["code"] == "codec_failed"
        disclosed = get(subject, f"/{run_id}/artifacts/{broken}/preview").json()["preview"]
        assert disclosed["kind"] == "codec_failed" and disclosed["body"]["code"] == "content_rejected"
        assert disclosed["coverage"]["covered"] == []
        # an original past the worker's input bound is never sent
        monkeypatch.setattr(run_artifacts, "CODEC_INPUT_BYTES", len(pdf) - 1)
        oversize = get(subject, f"/{run_id}/artifacts/{report}/pages/2")
        assert oversize.status_code == 413 and oversize.json()["code"] == "too_large"
        preview = get(subject, f"/{run_id}/artifacts/{report}/preview").json()["preview"]
        assert preview["kind"] == "codec_required" and preview["body"]["reason"] == "too_large"


def test_an_unnamed_or_unreachable_worker_is_disclosed_and_pages_answer_503(tmp_path):
    pdf = pdf_bytes(1)
    with owner_app(tmp_path / "unnamed") as subject:
        run_id, ref = started(subject, [("report", "application/pdf", pdf), ("draft", DOCX_MEDIA, docx_bytes())])
        preview = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 0)}/preview").json()["preview"]
        assert preview["kind"] == "codec_required" and preview["body"] == {"format": "pdf", "reason": "not_connected"}
        docx = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 1)}/preview").json()["preview"]
        assert docx["kind"] == "codec_required" and docx["body"]["format"] == "zip_container"
        page = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 0)}/pages/1")
        assert page.status_code == 503 and page.json()["code"] == "unavailable"
    missing = process_client(tmp_path / "no-such-worker.sock", secrets.token_bytes(32))
    with owner_app(tmp_path / "unreachable", document_worker=missing) as subject:
        run_id, ref = started(subject, [("report", "application/pdf", pdf)])
        preview = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 0)}/preview").json()["preview"]
        assert preview["kind"] == "codec_required" and preview["body"]["reason"] == "unavailable"
        assert preview["sha256"] is None and preview["coverage"]["covered"] == []
        image = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 0)}/pages/1/image")
        assert image.status_code == 503 and image.headers["content-type"].startswith("application/json")


class _DroppingService:
    """A worker that reads the request and then dies without an answer."""

    def serve_connection(self, connection, *, deadline):
        connection.read(deadline=deadline)
        connection.close()
        return "dropped"


def test_a_worker_that_drops_the_dialogue_is_unavailable_not_a_page(tmp_path):
    client, threads, _outcomes = in_thread_client(service=_DroppingService(), deadline_ms=5_000)
    with owner_app(tmp_path, document_worker=client) as subject:
        run_id, ref = started(subject, [("report", "application/pdf", pdf_bytes(1))])
        page = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 0)}/pages/1")
        assert page.status_code == 503
        preview = get(subject, f"/{run_id}/artifacts/{artifact_identity(ref, 0)}/preview").json()["preview"]
        assert preview["kind"] == "codec_required" and preview["body"]["reason"] == "unavailable"
    for thread in threads:
        thread.join(5)


def test_rendered_pages_are_cached_by_original_digest(tmp_path):
    client, threads, outcomes = in_thread_client()
    with owner_app(tmp_path, document_worker=client) as subject:
        run_id, ref = started(subject, [("report", "application/pdf", pdf_bytes(2))])
        path = f"/{run_id}/artifacts/{artifact_identity(ref, 0)}/pages/2"
        disclosed = get(subject, path).json()["preview"]
        image = get(subject, path + "/image")
        assert sha256(image.content).hexdigest() == disclosed["sha256"]
        assert get(subject, path + "/image").content == image.content
    for thread in threads:
        thread.join(5)
    assert outcomes == ["ok"]  # one render served the disclosure and both image reads


def test_the_supported_factory_accepts_only_a_named_worker_or_its_client(tmp_path):
    from app.workers.document_channel import DocumentWorkerConfiguration

    profile, _capability, arguments = configured(tmp_path)
    with pytest.raises(ValueError):
        create_app(tmp_path / "data", **arguments, document_worker=object())
    with pytest.raises(ValueError):  # only the verified cp-document pair root is an endpoint
        create_app(tmp_path / "data", **arguments, document_worker=DocumentWorkerConfiguration(
            pair_root=str(tmp_path / "cp-document"), requester_boot_id="control-boot"))
    app = create_app(tmp_path / "data", **arguments, document_worker=DocumentWorkerConfiguration(
        pair_root="/run/deeptwin/ipc/cp-document", requester_boot_id="control-boot"))
    assert type(app.state.first_party_exports["run-artifacts.service"]._codec) is DocumentCodecClient


def test_the_index_lists_every_runs_artifacts_in_bounded_filtered_pages(tmp_path):
    with owner_app(tmp_path) as subject:
        first_run, first = started(subject, [("report", "text/plain", b"one"), ("chart", "image/png", PNG)])
        second_run, second = started(subject, [("paper", "application/pdf", pdf_bytes(1)),
                                               ("table", "text/csv", b"a,b\n1,2\n")])
        whole = index(subject)
        assert whole.status_code == 200, whole.text
        value = whole.json()
        assert [(item["run_id"], item["artifact_id"]) for item in value["artifacts"]] == [
            (first_run, artifact_identity(first, 0)), (first_run, artifact_identity(first, 1)),
            (second_run, artifact_identity(second, 0)), (second_run, artifact_identity(second, 1))]
        assert value["total"] == 4 and value["next_cursor"] is None
        assert value["runs"] == {"scanned": 2, "unreadable": 0, "omitted": 0}
        assert value["artifacts"][0]["availability"] == "available"
        paged = index(subject, "?limit=3").json()
        assert len(paged["artifacts"]) == 3 and paged["next_cursor"] == "3"
        rest = index(subject, "?limit=3&cursor=3").json()
        assert [item["role"] for item in rest["artifacts"]] == ["table"] and rest["next_cursor"] is None
        assert [item["role"] for item in index(subject, "?media_type=application%2Fpdf").json()["artifacts"]] == ["paper"]
        assert [item["role"] for item in index(subject, "?media_type=text/*").json()["artifacts"]] == ["report", "table"]
        only = index(subject, f"?run_id={second_run}").json()
        assert {item["run_id"] for item in only["artifacts"]} == {second_run}
        # each listed artifact opens through its run's own routes
        entry = value["artifacts"][2]
        assert get(subject, f"/{entry['run_id']}/artifacts/{entry['artifact_id']}").json()["role"] == "paper"
        assert index(subject, f"?run_id={uuid4()}").status_code == 404
        for bad in ("?limit=0", "?limit=101", "?cursor=01", "?cursor=-1", "?media_type=bad",
                    "?run_id=nope", "?other=1", "?limit=1&limit=2"):
            response = index(subject, bad)
            assert response.status_code == 400, bad
        subject.client.cookies.clear()
        assert index(subject).status_code == 401


def test_an_empty_vault_indexes_nothing(tmp_path):
    with owner_app(tmp_path) as subject:
        value = index(subject).json()
        assert value["artifacts"] == [] and value["total"] == 0
        assert value["runs"] == {"scanned": 0, "unreadable": 0, "omitted": 0}
        head = subject.client.head(subject.profile.base_path + "api/v1/artifacts", headers=headers(subject.profile))
        assert head.status_code == 200 and head.content == b""
